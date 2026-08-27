from __future__ import annotations

import hashlib
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterator

from openpyxl import load_workbook

from .models import Analytics, HierarchyNode, JournalEntry
from .normalization import account_key, money, original_text, stable_hash, text_key


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _rows(path: Path) -> Iterator[tuple[str, int, dict[str, Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        iterator = sheet.iter_rows(values_only=True)
        try:
            headers = [original_text(v) for v in next(iterator)]
        except StopIteration:
            continue
        for index, values in enumerate(iterator, 2):
            if any(value is not None for value in values):
                yield sheet.title, index, dict(zip(headers, values, strict=False))
    workbook.close()


def _value(row: dict[str, Any], *names: str) -> Any:
    keyed = {text_key(k): v for k, v in row.items()}
    for name in names:
        if text_key(name) in keyed:
            return keyed[text_key(name)]
    return None


def _date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.fromisoformat(str(value)).date()


def read_journals(
    paths: list[Path],
    organization_lookup: dict[str, str] | None = None,
    department_lookup: dict[str, str] | None = None,
) -> list[JournalEntry]:
    entries: list[JournalEntry] = []
    organization_lookup = organization_lookup or {}
    department_lookup = department_lookup or {}
    for path in sorted(paths):
        sha = file_sha256(path)
        for sheet, row_number, row in _rows(path):
            org_original = original_text(_value(row, "Организация"))
            org = organization_lookup.get(text_key(org_original), org_original)

            def analytics(side: str) -> Analytics:
                dept_original = original_text(_value(row, f"ЦФО (подразделение) {side}"))
                dept = department_lookup.get(text_key(dept_original), dept_original)
                return Analytics(
                    org,
                    dept,
                    original_text(_value(row, f"Направление деятельности {side}")),
                    tuple(
                        original_text(_value(row, name))
                        for name in (
                            f"Субконто {side}",
                            f"Субконто2 {side}",
                            f"Субконто3 {side}",
                            f"Субконто {side.lower()}4",
                        )
                    ),
                )

            source_id = stable_hash([sha, sheet, row_number, _value(row, "№")])
            entries.append(
                JournalEntry(
                    sha,
                    path.name,
                    sheet,
                    row_number,
                    source_id,
                    _date(_value(row, "Дата")),
                    original_text(_value(row, "Документ")),
                    org,
                    org_original,
                    original_text(_value(row, "Сценарий")),
                    account_key(_value(row, "Счет Дт")),
                    analytics("Дт"),
                    original_text(_value(row, "Валюта Дт")),
                    account_key(_value(row, "Счет Кт")),
                    analytics("Кт"),
                    original_text(_value(row, "Валюта Кт")),
                    money(_value(row, "Сумма в валюте учета")),
                    money(_value(row, "Сумма в валюте отчетности"))
                    if _value(row, "Сумма в валюте отчетности") not in (None, "")
                    else None,
                    original_text(_value(row, "Содержание")),
                    original_text(_value(row, "ИдентификаторФинЗаписи")) or None,
                    row,
                )
            )
    return entries


def read_hierarchy(path: Path) -> list[HierarchyNode]:
    nodes: list[HierarchyNode] = []
    for _, _, row in _rows(path):
        code = original_text(_value(row, "Код", "node_code"))
        name = original_text(
            _value(
                row,
                "Наименование",
                "ЦФО",
                "Сокращенное юр. наименование",
                "Полное юр. наименование",
                "Головная организация",
            )
        )
        node_type = original_text(_value(row, "Тип", "node_type")).upper()
        owner = (
            original_text(_value(row, "Код организации-владельца", "owner_organization_code"))
            or None
        )
        if not node_type:
            node_type = (
                "DEPARTMENT"
                if owner
                else ("ROOT" if not _value(row, "Код родителя", "parent_code") else "ORGANIZATION")
            )
        nodes.append(
            HierarchyNode(
                code,
                name,
                node_type,
                original_text(_value(row, "Код родителя", "parent_code")) or None,
                original_text(_value(row, "Код корня", "root_code")) or None,
                owner,
                original_text(_value(row, "Верхний уровень иерархии")) or None,
                original_text(_value(row, "Код эквивалентности", "equivalence_code")) or None,
                text_key(_value(row, "Без перевыставления")) in {"да", "true", "1"},
            )
        )
    return nodes
