from __future__ import annotations

import json
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from .models import AnalysisResult, Correction, serializable

EXPORT_COLUMNS = [
    "СчетДт",
    "СчетКт",
    "ВалютаДт",
    "ВалютаКт",
    "ВидОперации",
    "ПодразделениеДт",
    "ПодразделениеКт",
    "НаправлениеДеятельностиДт",
    "НаправлениеДеятельностиКт",
    "СуммаВВалютеУчета",
    "СуммаВВалютеОтчетности",
    "СуммаВВалютеДт",
    "СуммаВВалютеКт",
    "КоличествоДт",
    "КоличествоКт",
    "Содержание",
    "СчетДтИсточник",
    "СчетКтИсточник",
    "ИдентификаторФинЗаписи",
    "ПравилоДт",
    "ПравилоКт",
    "СубконтоДт1",
    "СубконтоДт2",
    "СубконтоДт3",
    "СубконтоКт1",
    "СубконтоКт2",
    "СубконтоКт3",
]

CONTROL_SHEETS = [
    "Итоги",
    "Параметры_запуска",
    "Область_иерархии",
    "Выбранные_счета",
    "Пары_организаций",
    "Пары_подразделений",
    "Обороты_по_счетам",
    "Остатки_аналитик",
    "Кандидаты",
    "Готовые_проводки",
    "Блокировки",
    "Контроль_до_после",
    "Исходные_строки",
    "Исключенные_строки",
    "Проверка_экспорта",
]


def write_json(path: Path, value: Any) -> None:
    path.write_text(
        json.dumps(serializable(value), ensure_ascii=False, indent=2, sort_keys=True),
        encoding="utf-8",
    )


def write_jsonl(path: Path, values: list[Any]) -> None:
    path.write_text(
        "".join(
            json.dumps(serializable(v), ensure_ascii=False, sort_keys=True) + "\n" for v in values
        ),
        encoding="utf-8",
    )


def write_control(
    path: Path, summary: dict[str, Any], result: AnalysisResult, selection: dict[str, Any]
) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for name in CONTROL_SHEETS:
        wb.create_sheet(name)
    for key, value in summary.items():
        wb["Итоги"].append(
            [
                key,
                json.dumps(serializable(value), ensure_ascii=False)
                if isinstance(value, (list, dict))
                else value,
            ]
        )
    for key, value in selection.items():
        wb["Параметры_запуска"].append([key, str(value)])
    wb["Остатки_аналитик"].append(
        ["Период", "Счет", "Организация", "Подразделение", "Дт", "Кт", "Остаток"]
    )
    for b in result.buckets:
        wb["Остатки_аналитик"].append(
            [
                b.period,
                b.account,
                b.analytics.organization_code,
                b.analytics.department_code,
                b.debit_turnover,
                b.credit_turnover,
                b.net,
            ]
        )
    wb["Готовые_проводки"].append(["correction_id", "status", "Счет", "Сумма", "Дт", "Кт"])
    for c in result.corrections:
        wb["Готовые_проводки"].append(
            [
                c.correction_id,
                c.status,
                c.account,
                c.amount,
                c.target.department_code,
                c.source.department_code,
            ]
        )
    wb["Блокировки"].append(
        ["Код", "Описание", "Период", "Счет", "Подразделение", "Сумма", "Строки"]
    )
    for b in result.blockages:
        wb["Блокировки"].append(
            [
                b.code,
                b.description,
                b.period,
                b.account,
                b.analytics.department_code,
                b.amount,
                "; ".join(b.row_refs),
            ]
        )
    wb["Исключенные_строки"].append(["source_row_id", "Причина"])
    for row in result.excluded:
        wb["Исключенные_строки"].append([row.get("source_row_id"), row.get("reason")])
    wb.save(path)


def export_rows(
    corrections: list[Correction], path: Path, rules: dict[str, Any], run_id: str, rules_sha: str
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Загрузка_A_AA"
    ws.append(EXPORT_COLUMNS)
    rule_code = rules.get("export", {}).get("rule_code", "HAT_TRANSFER_EXACT_V1")
    operation = rules.get("export", {}).get("operation_kind", "REPOST")
    for c in corrections:
        content = (
            f"run_id={run_id}; correction_id={c.correction_id}; account={c.account}; "
            f"source={c.source.organization_code}/{c.source.department_code}; "
            f"target={c.target.organization_code}/{c.target.department_code}; "
            f"before={c.source_net_before}/{c.target_net_before}; after={c.source_net_after}/{c.target_net_after}; "
            f"source_refs={','.join(c.source_refs)}; target_refs={','.join(c.target_refs)}; "
            f"pair={c.pair_mapping_id}; rules_sha256={rules_sha}"
        )
        ws.append(
            [
                c.account,
                c.account,
                c.currency,
                c.currency,
                operation,
                c.target.department_code,
                c.source.department_code,
                c.target.activity,
                c.source.activity,
                c.amount,
                None,
                c.amount,
                c.amount,
                None,
                None,
                content,
                c.account,
                c.account,
                c.correction_id,
                f"{rule_code}:{c.account}",
                f"{rule_code}:{c.account}",
                *c.target.subconto[:3],
                *c.source.subconto[:3],
            ]
        )
    wb.save(path)
    check = load_workbook(path, read_only=True, data_only=True)
    ws2 = check["Загрузка_A_AA"]
    headers = [cell.value for cell in next(ws2.iter_rows())]
    count = sum(1 for _ in ws2.iter_rows())
    check.close()
    if headers != EXPORT_COLUMNS or count != len(corrections):
        path.unlink(missing_ok=True)
        raise ValueError("BLOCKED_OUTPUT_ROUNDTRIP_MISMATCH")


def summary_for(
    run_id: str,
    root_code: str,
    accounts: set[str],
    scenario: str,
    scope: set[str],
    result: AnalysisResult,
) -> dict[str, Any]:
    ready = [c for c in result.corrections if c.status == "READY"]
    amounts: dict[str, str] = {}
    for account in sorted(accounts):
        amounts[account] = str(sum((c.amount for c in ready if c.account == account), start=0))
    return {
        "run_id": run_id,
        "root_organization_code": root_code,
        "scope_node_count": len(scope),
        "selected_accounts": sorted(accounts),
        "periods": sorted({b.period for b in result.buckets}),
        "scenario": scenario,
        "source_row_count": len({e.source_row_id for e in result.events}),
        "account_event_count": len(result.events),
        "ready_count": len(ready),
        "blocked_count": len(result.blockages),
        "review_count": 0,
        "excluded_count": len(result.excluded),
        "ready_amount_by_account": amounts,
        "blocked_amount_by_account": dict(Counter(b.account for b in result.blockages)),
        "run_status": "READY"
        if ready and not result.blockages
        else ("BLOCKED" if result.blockages else "NO_ACTION"),
    }
