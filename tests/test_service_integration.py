import json
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook

from hierarchy_account_transfer.artifacts import CONTROL_SHEETS, EXPORT_COLUMNS
from hierarchy_account_transfer.service import HatService


def workbook(path: Path, headers: list[str], rows: list[list[object]]) -> None:
    book = Workbook()
    sheet = book.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    book.save(path)


def test_full_inspect_analyze_export(tmp_path: Path):
    hierarchy = tmp_path / "hierarchy.xlsx"
    workbook(
        hierarchy,
        ["Код", "Наименование", "Тип", "Код родителя", "Код организации-владельца"],
        [
            ["ROOT", "Root", "ROOT", None, None],
            ["ORG", "Org", "ORGANIZATION", "ROOT", None],
            ["S", "Source", "DEPARTMENT", "ORG", "ORG"],
            ["T", "Target", "DEPARTMENT", "ORG", "ORG"],
        ],
    )
    journal = tmp_path / "journal.xlsx"
    headers = [
        "№",
        "Дата",
        "Организация",
        "Счет Дт",
        "Валюта Дт",
        "ЦФО (подразделение) Дт",
        "Направление деятельности Дт",
        "Субконто Дт",
        "Субконто2 Дт",
        "Субконто3 Дт",
        "Субконто дт4",
        "Счет Кт",
        "Валюта Кт",
        "ЦФО (подразделение) Кт",
        "Направление деятельности Кт",
        "Субконто Кт",
        "Субконто2 Кт",
        "Субконто3 Кт",
        "Субконто кт4",
        "Сумма в валюте учета",
        "Сценарий",
    ]
    common = ["Activity", "1", "2", "3", "4"]
    workbook(
        journal,
        headers,
        [
            [
                1,
                "2025-01-15",
                "Org",
                "A",
                "RUB",
                "Source",
                *common,
                "X",
                "RUB",
                "Target",
                *common,
                1000,
                "Actual",
            ],
            [
                2,
                "2025-01-15",
                "Org",
                "X",
                "RUB",
                "Source",
                *common,
                "A",
                "RUB",
                "Target",
                *common,
                1000,
                "Actual",
            ],
        ],
    )
    rules_path = tmp_path / "rules.yaml"
    rules_path.write_text(
        yaml.safe_dump(
            {
                "contract_version": "HAT-TRANSFER-CONTRACT-v1",
                "calculation": {"amount_tolerance": "0.05", "allow_partial_transfer": False},
                "selection": {"include_subaccounts": False},
                "pairing": {
                    "department_pairs": [
                        {
                            "id": "pair",
                            "source_organization_code": "ORG",
                            "target_organization_code": "ORG",
                            "source_department_code": "S",
                            "target_department_code": "T",
                        }
                    ]
                },
                "export": {
                    "organization_encoding_mode": "DOCUMENT_LEVEL",
                    "all_or_nothing_export": True,
                    "document_date_policy": "TARGET_CREDIT_DATE_IF_UNIQUE_ELSE_PERIOD_END",
                },
            }
        ),
        encoding="utf-8",
    )
    service = HatService(tmp_path / "runs")
    inspection = service.inspect([journal], hierarchy)
    assert inspection["available_roots"][0]["included_node_count"] == 4
    run_id, summary = service.analyze(
        inspection["inspection_id"],
        "ROOT",
        {"A"},
        "Actual",
        __import__("datetime").date(2025, 1, 1),
        __import__("datetime").date(2025, 1, 31),
        rules_path,
    )
    assert summary["ready_count"] == 1
    control = load_workbook(tmp_path / "runs" / run_id / "run_control.xlsx", read_only=True)
    assert control.sheetnames == CONTROL_SHEETS
    control.close()
    manifest = service.export(run_id)
    assert manifest["status"] == "EXPORTED"
    exported = load_workbook(
        tmp_path / "runs" / run_id / "export" / manifest["files"][0], read_only=True
    )
    assert exported.sheetnames == ["Загрузка_A_AA"]
    assert [cell.value for cell in next(exported.active.iter_rows())] == EXPORT_COLUMNS
    assert exported.active.max_row == 2
    exported.close()
    assert json.loads((tmp_path / "runs" / run_id / "summary.json").read_text())["ready_count"] == 1
